from __future__ import annotations

import calendar
import os
import re
import sys
import unicodedata
from dataclasses import dataclass
from datetime import date
from pathlib import Path
from typing import Any

import requests
from dotenv import load_dotenv
from openpyxl import load_workbook


MONTHS = {
    "ene": 1,
    "enero": 1,
    "feb": 2,
    "febrero": 2,
    "mar": 3,
    "marzo": 3,
    "abr": 4,
    "abril": 4,
    "may": 5,
    "mayo": 5,
    "jun": 6,
    "junio": 6,
    "jul": 7,
    "julio": 7,
    "ago": 8,
    "agos": 8,
    "agosto": 8,
    "sep": 9,
    "sept": 9,
    "septiembre": 9,
    "oct": 10,
    "octubre": 10,
    "nov": 11,
    "noviembre": 11,
    "dic": 12,
    "diciembre": 12,
}

MONTH_NAMES = {
    1: "Enero",
    2: "Febrero",
    3: "Marzo",
    4: "Abril",
    5: "Mayo",
    6: "Junio",
    7: "Julio",
    8: "Agosto",
    9: "Septiembre",
    10: "Octubre",
    11: "Noviembre",
    12: "Diciembre",
}

HEADERS_ROW = 3


@dataclass(frozen=True)
class WorkItem:
    row_number: int
    sync_key: str
    lineamiento: str
    plan_otic: str
    entregable: str
    descripcion: str
    actividad: str
    code: str
    responsable: str
    tareas_relacionadas: str
    mes: str
    producto: str
    estado: str
    link: str
    month_numbers: tuple[int, ...]
    month_statuses: tuple[str, ...]


class GitHub:
    def __init__(self, token: str) -> None:
        self.session = requests.Session()
        self.session.headers.update(
            {
                "Authorization": f"Bearer {token}",
                "Accept": "application/vnd.github+json",
                "X-GitHub-Api-Version": "2022-11-28",
            }
        )

    def rest(self, method: str, url: str, **kwargs: Any) -> Any:
        response = self.session.request(method, url, **kwargs)
        if response.status_code >= 400:
            raise RuntimeError(f"{method} {url} failed: {response.status_code} {response.text}")
        if not response.text:
            return None
        return response.json()

    def graphql(self, query: str, variables: dict[str, Any] | None = None) -> Any:
        response = self.session.post(
            "https://api.github.com/graphql",
            json={"query": query, "variables": variables or {}},
            headers={"Accept": "application/vnd.github+json"},
        )
        if response.status_code >= 400:
            raise RuntimeError(f"GraphQL failed: {response.status_code} {response.text}")
        payload = response.json()
        if payload.get("errors"):
            raise RuntimeError(f"GraphQL errors: {payload['errors']}")
        return payload["data"]


def clean(value: Any) -> str:
    if value is None:
        return ""
    return re.sub(r"[ \t]+\n", "\n", str(value).strip())


def slug(value: str) -> str:
    normalized = unicodedata.normalize("NFKD", value).encode("ascii", "ignore").decode("ascii")
    return re.sub(r"[^a-z0-9]+", "_", normalized.lower()).strip("_")


def split_roles(value: str) -> list[str]:
    value = value.replace("Cientifico", "Científico")
    value = value.replace("datos jr", "Datos Jr")
    parts = re.split(r"\s*/\s*|\s+y\s+", value)
    return [p.strip() for p in parts if p.strip()]


def parse_months(value: str) -> tuple[int, ...]:
    if not value:
        return ()
    tokens = [t for t in re.split(r"[/,\-\s]+", value.lower()) if t]
    found: list[int] = []
    for token in tokens:
        normalized = unicodedata.normalize("NFKD", token).encode("ascii", "ignore").decode("ascii")
        month = MONTHS.get(normalized)
        if month and month not in found:
            found.append(month)
    return tuple(found)


def full_month_range(months: tuple[int, ...], year: int) -> tuple[str | None, str | None]:
    if not months:
        return None, None
    start_month = min(months)
    end_month = max(months)
    _, last_day = calendar.monthrange(year, end_month)
    return date(year, start_month, 1).isoformat(), date(year, end_month, last_day).isoformat()


def load_workplan(path: Path) -> list[WorkItem]:
    workbook = load_workbook(path, data_only=True)
    worksheet = workbook.active
    headers = [clean(cell.value) for cell in worksheet[HEADERS_ROW]]
    index = {header: position for position, header in enumerate(headers)}

    required = [
        "Lineamiento MGGTI",
        "PLAN OTIC 2026",
        "Entregable (Soporte)",
        "Descripción",
        "Actividad",
        "ID",
        "Responsable",
        "Tareas relacionadas",
        "Mes",
        "Relación de producto",
        "ESTADO",
        "LINK",
    ]
    missing = [header for header in required if header not in index]
    if missing:
        raise RuntimeError(f"Missing required Excel columns: {missing}")

    raw_items: list[dict[str, Any]] = []
    for row_number, row in enumerate(worksheet.iter_rows(min_row=HEADERS_ROW + 1, values_only=True), start=HEADERS_ROW + 1):
        if not any(row):
            continue
        code = clean(row[index["ID"]])
        actividad = clean(row[index["Actividad"]])
        if not code and not actividad:
            continue
        mes = clean(row[index["Mes"]])
        month_statuses = tuple(
            clean(row[index[name]])
            for name in MONTH_NAMES.values()
            if name in index and clean(row[index[name]])
        )
        raw_items.append(
            {
                "row_number": row_number,
                "lineamiento": clean(row[index["Lineamiento MGGTI"]]),
                "plan_otic": clean(row[index["PLAN OTIC 2026"]]),
                "entregable": clean(row[index["Entregable (Soporte)"]]),
                "descripcion": clean(row[index["Descripción"]]),
                "actividad": actividad,
                "code": code,
                "responsable": clean(row[index["Responsable"]]),
                "tareas_relacionadas": clean(row[index["Tareas relacionadas"]]),
                "mes": mes,
                "producto": clean(row[index["Relación de producto"]]) or "Sin producto asociado",
                "estado": clean(row[index["ESTADO"]]),
                "link": clean(row[index["LINK"]]),
                "month_numbers": parse_months(mes),
                "month_statuses": month_statuses,
            }
        )
    code_counts: dict[str, int] = {}
    items: list[WorkItem] = []
    for raw in raw_items:
        code = raw["code"]
        code_counts[code] = code_counts.get(code, 0) + 1
        sync_key = code if code_counts[code] == 1 else f"{code}-{code_counts[code]}"
        items.append(WorkItem(sync_key=sync_key, **raw))
    return items


def get_repo(gh: GitHub, repository: str) -> dict[str, Any]:
    return gh.rest("GET", f"https://api.github.com/repos/{repository}")


def ensure_project(gh: GitHub, owner_login: str, repo_node_id: str, title: str) -> dict[str, str]:
    data = gh.graphql(
        """
        query($login: String!, $first: Int!) {
          user(login: $login) {
            id
            projectsV2(first: $first) {
              nodes { id number title url }
            }
          }
        }
        """,
        {"login": owner_login, "first": 100},
    )
    owner = data["user"]
    for project in owner["projectsV2"]["nodes"]:
        if project["title"] == title:
            link_project_to_repo(gh, project["id"], repo_node_id)
            return project

    created = gh.graphql(
        """
        mutation($ownerId: ID!, $title: String!) {
          createProjectV2(input: {ownerId: $ownerId, title: $title}) {
            projectV2 { id number title url }
          }
        }
        """,
        {"ownerId": owner["id"], "title": title},
    )["createProjectV2"]["projectV2"]
    link_project_to_repo(gh, created["id"], repo_node_id)
    return created


def link_project_to_repo(gh: GitHub, project_id: str, repo_node_id: str) -> None:
    try:
        gh.graphql(
            """
            mutation($projectId: ID!, $repositoryId: ID!) {
              linkProjectV2ToRepository(input: {projectId: $projectId, repositoryId: $repositoryId}) {
                repository { id }
              }
            }
            """,
            {"projectId": project_id, "repositoryId": repo_node_id},
        )
    except RuntimeError as exc:
        message = str(exc)
        if "already" not in message.lower() and "exists" not in message.lower():
            raise


def get_project_fields(gh: GitHub, project_id: str) -> dict[str, dict[str, Any]]:
    nodes = gh.graphql(
        """
        query($projectId: ID!) {
          node(id: $projectId) {
            ... on ProjectV2 {
              fields(first: 50) {
                nodes {
                  ... on ProjectV2FieldCommon { id name dataType }
                  ... on ProjectV2SingleSelectField {
                    id
                    name
                    dataType
                    options { id name }
                  }
                }
              }
            }
          }
        }
        """,
        {"projectId": project_id},
    )["node"]["fields"]["nodes"]
    return {field["name"]: field for field in nodes if field}


def ensure_field(
    gh: GitHub,
    project_id: str,
    fields: dict[str, dict[str, Any]],
    name: str,
    data_type: str,
    options: list[str] | None = None,
) -> dict[str, Any]:
    if name in fields:
        return fields[name]
    payload: dict[str, Any] = {"projectId": project_id, "name": name, "dataType": data_type}
    if options:
        payload["singleSelectOptions"] = [
            {"name": option, "color": "GRAY", "description": option} for option in options
        ]
    field = gh.graphql(
        """
        mutation($input: CreateProjectV2FieldInput!) {
          createProjectV2Field(input: $input) {
            projectV2Field {
              ... on ProjectV2FieldCommon { id name dataType }
              ... on ProjectV2SingleSelectField {
                id
                name
                dataType
                options { id name }
              }
            }
          }
        }
        """,
        {"input": payload},
    )["createProjectV2Field"]["projectV2Field"]
    fields[name] = field
    return field


def ensure_labels(gh: GitHub, repository: str, labels: set[str]) -> None:
    existing = {item["name"] for item in gh.rest("GET", f"https://api.github.com/repos/{repository}/labels?per_page=100")}
    for label in sorted(labels - existing):
        color = "5319e7" if label.startswith("rol:") else "1d76db"
        gh.rest(
            "POST",
            f"https://api.github.com/repos/{repository}/labels",
            json={"name": label, "color": color, "description": "Creado por sincronizacion del plan de trabajo"},
        )


def list_milestones(gh: GitHub, repository: str) -> dict[str, int]:
    milestones: dict[str, int] = {}
    for state in ("open", "closed"):
        items = gh.rest(
            "GET",
            f"https://api.github.com/repos/{repository}/milestones?state={state}&per_page=100",
        )
        milestones.update({item["title"]: item["number"] for item in items})
    return milestones


def milestone_title(product: str) -> str:
    return re.sub(r"\s+", " ", product).strip()[:250] or "Sin producto asociado"


def ensure_milestones(gh: GitHub, repository: str, items: list[WorkItem], year: int) -> dict[str, int]:
    milestones = list_milestones(gh, repository)
    for product in sorted({item.producto for item in items}):
        title = milestone_title(product)
        if title in milestones:
            milestones[product] = milestones[title]
            continue
        product_items = [item for item in items if item.producto == product]
        months = tuple(sorted({month for item in product_items for month in item.month_numbers}))
        _, due_on = full_month_range(months, year)
        payload: dict[str, Any] = {
            "title": title,
            "description": "Producto asociado al plan de trabajo importado desde Excel.",
        }
        if due_on:
            payload["due_on"] = f"{due_on}T23:59:59Z"
        try:
            milestone = gh.rest("POST", f"https://api.github.com/repos/{repository}/milestones", json=payload)
            milestones[title] = milestone["number"]
            milestones[product] = milestone["number"]
        except RuntimeError as exc:
            if "already_exists" not in str(exc):
                raise
            milestones = list_milestones(gh, repository)
            if title not in milestones:
                raise
            milestones[product] = milestones[title]
    return milestones


def issue_title(item: WorkItem) -> str:
    prefix = f"[{item.sync_key}] " if item.sync_key else ""
    return f"{prefix}{item.actividad}"[:250]


def issue_body(item: WorkItem, year: int) -> str:
    start, end = full_month_range(item.month_numbers, year)
    roles = ", ".join(split_roles(item.responsable)) or item.responsable
    sprint = ", ".join(f"Sprint {month:02d} - {MONTH_NAMES[month]} {year}" for month in item.month_numbers)
    lines = [
        "## Contexto",
        f"- Lineamiento MGGTI: {item.lineamiento or 'Sin dato'}",
        f"- Plan OTIC 2026: {item.plan_otic or 'Sin dato'}",
        f"- Entregable: {item.entregable or 'Sin dato'}",
        f"- Producto/Milestone: {item.producto}",
        f"- Responsable/Rol: {roles or 'Sin dato'}",
        f"- Mes original: {item.mes or 'Sin dato'}",
        f"- Sprint mensual: {sprint or 'Sin sprint asignado'}",
        f"- Ventana: {start or 'Sin fecha'} a {end or 'Sin fecha'}",
        "",
        "## Descripcion",
        item.descripcion or "Sin descripcion.",
        "",
        "## Actividad",
        item.actividad or "Sin actividad.",
    ]
    if item.tareas_relacionadas:
        lines.extend(["", "## Tareas relacionadas", item.tareas_relacionadas])
    if item.estado:
        lines.extend(["", "## Estado inicial desde Excel", item.estado])
    if item.link:
        lines.extend(["", "## Link / soporte", item.link])
    lines.extend(
        [
            "",
            "<!-- workplan-sync -->",
            f"workplan_id: {item.code}",
            f"workplan_key: {item.sync_key}",
            f"excel_row: {item.row_number}",
        ]
    )
    return "\n".join(lines)


def status_from_item(item: WorkItem, today: date, year: int) -> str:
    if not item.month_numbers:
        return "Todo"
    start_month = min(item.month_numbers)
    end_month = max(item.month_numbers)
    if year < today.year or (year == today.year and end_month < today.month):
        return "Done"
    if year == today.year and start_month <= today.month <= end_month:
        return "In Progress"
    return "Todo"


def find_existing_issues(gh: GitHub, repository: str) -> dict[str, dict[str, Any]]:
    issues: dict[str, dict[str, Any]] = {}
    page = 1
    while True:
        page_items = gh.rest(
            "GET",
            f"https://api.github.com/repos/{repository}/issues?state=all&per_page=100&page={page}",
        )
        if not page_items:
            break
        for issue in page_items:
            if "pull_request" in issue:
                continue
            match = re.search(r"\[([A-Za-z]+-\d+(?:-\d+)?)\]", issue["title"])
            if match:
                issues[match.group(1)] = issue
        page += 1
    return issues


def assignees_for(item: WorkItem) -> list[str]:
    assignees: list[str] = []
    for role in split_roles(item.responsable):
        env_name = f"ROLE_{slug(role).upper()}"
        login = os.getenv(env_name)
        if login and login not in assignees:
            assignees.append(login)
    return assignees


def create_or_update_issue(
    gh: GitHub,
    repository: str,
    item: WorkItem,
    year: int,
    milestones: dict[str, int],
    existing: dict[str, dict[str, Any]],
) -> dict[str, Any]:
    labels = [f"rol: {role}" for role in split_roles(item.responsable)]
    labels += [f"sprint: {month:02d}-{MONTH_NAMES[month].lower()}" for month in item.month_numbers]
    labels.append("workplan")
    payload = {
        "title": issue_title(item),
        "body": issue_body(item, year),
        "labels": labels,
        "milestone": milestones.get(item.producto),
        "assignees": assignees_for(item),
    }
    if item.sync_key in existing:
        issue = existing[item.sync_key]
        try:
            return gh.rest("PATCH", issue["url"], json=payload)
        except RuntimeError as exc:
            if payload["assignees"] and "assignee" in str(exc).lower():
                payload["assignees"] = []
                return gh.rest("PATCH", issue["url"], json=payload)
            raise
    try:
        issue = gh.rest("POST", f"https://api.github.com/repos/{repository}/issues", json=payload)
    except RuntimeError as exc:
        if payload["assignees"] and "assignee" in str(exc).lower():
            payload["assignees"] = []
            issue = gh.rest("POST", f"https://api.github.com/repos/{repository}/issues", json=payload)
        else:
            raise
    existing[item.sync_key] = issue
    return issue


def add_issue_to_project(gh: GitHub, project_id: str, issue_node_id: str) -> str:
    data = gh.graphql(
        """
        mutation($projectId: ID!, $contentId: ID!) {
          addProjectV2ItemById(input: {projectId: $projectId, contentId: $contentId}) {
            item { id }
          }
        }
        """,
        {"projectId": project_id, "contentId": issue_node_id},
    )
    return data["addProjectV2ItemById"]["item"]["id"]


def set_project_field(gh: GitHub, project_id: str, item_id: str, field_id: str, value: dict[str, Any]) -> None:
    gh.graphql(
        """
        mutation($projectId: ID!, $itemId: ID!, $fieldId: ID!, $value: ProjectV2FieldValue!) {
          updateProjectV2ItemFieldValue(
            input: {projectId: $projectId, itemId: $itemId, fieldId: $fieldId, value: $value}
          ) { projectV2Item { id } }
        }
        """,
        {"projectId": project_id, "itemId": item_id, "fieldId": field_id, "value": value},
    )


def single_select_option_id(field: dict[str, Any], name: str) -> str | None:
    for option in field.get("options", []):
        if option["name"] == name:
            return option["id"]
    return None


def get_project_views(gh: GitHub, project_id: str) -> dict[str, dict[str, Any]]:
    nodes = gh.graphql(
        """
        query($projectId: ID!) {
          node(id: $projectId) {
            ... on ProjectV2 {
              views(first: 50) {
                nodes { id name number layout filter }
              }
            }
          }
        }
        """,
        {"projectId": project_id},
    )["node"]["views"]["nodes"]
    return {view["name"]: view for view in nodes if view}


def ensure_view(
    gh: GitHub,
    project_id: str,
    views: dict[str, dict[str, Any]],
    name: str,
    layout: str,
    visible_field_ids: list[str] | None,
    filter_query: str = "",
) -> dict[str, Any]:
    configuration = {"visibleFieldIds": visible_field_ids} if visible_field_ids is not None else None
    if name in views:
        input_payload: dict[str, Any] = {
            "viewId": views[name]["id"],
            "name": name,
            "layout": layout,
            "filter": filter_query,
        }
        if configuration is not None:
            input_payload["configuration"] = configuration
        data = gh.graphql(
            """
            mutation($input: UpdateProjectV2ViewInput!) {
              updateProjectV2View(input: $input) {
                projectV2View { id name number layout filter }
              }
            }
            """,
            {"input": input_payload},
        )
        views[name] = data["updateProjectV2View"]["projectV2View"]
        return views[name]

    input_payload = {"projectId": project_id, "name": name, "layout": layout}
    if configuration is not None:
        input_payload["configuration"] = configuration
    data = gh.graphql(
        """
        mutation($input: CreateProjectV2ViewInput!) {
          createProjectV2View(input: $input) {
            projectV2View { id name number layout filter }
          }
        }
        """,
        {"input": input_payload},
    )
    view = data["createProjectV2View"]["projectV2View"]
    views[name] = view
    if filter_query:
        return ensure_view(gh, project_id, views, name, layout, visible_field_ids, filter_query)
    return view


def ensure_standard_views(gh: GitHub, project_id: str, fields: dict[str, dict[str, Any]]) -> None:
    views = get_project_views(gh, project_id)

    def ids(*names: str) -> list[str]:
        return [fields[name]["id"] for name in names if name in fields and fields[name].get("id")]

    common = ids("Title", "Assignees", "Status", "Rol", "Sprint", "Milestone", "Inicio", "Fin", "Labels")
    ensure_view(gh, project_id, views, "Backlog", "BOARD_LAYOUT", common, "-status:Done")
    ensure_view(gh, project_id, views, "Priority board", "BOARD_LAYOUT", common, "-status:Done")
    ensure_view(gh, project_id, views, "Team items", "TABLE_LAYOUT", common, "")
    ensure_view(gh, project_id, views, "Roadmap", "ROADMAP_LAYOUT", None, "")
    ensure_view(gh, project_id, views, "My items", "TABLE_LAYOUT", common, "assignee:@me")
    ensure_view(gh, project_id, views, "Milestones", "TABLE_LAYOUT", ids("Title", "Milestone", "Status", "Sprint", "Inicio", "Fin", "Rol"), "")


def main() -> int:
    load_dotenv()
    token = os.getenv("GITHUB_TOKEN")
    repository = os.getenv("GITHUB_REPOSITORY", "morenoluisg/template_mads_project")
    owner = os.getenv("GITHUB_PROJECT_OWNER", repository.split("/")[0])
    project_title = os.getenv("GITHUB_PROJECT_TITLE", repository.split("/")[1])
    excel_path = Path(os.getenv("WORKPLAN_EXCEL_PATH", "data/raw/Plan de Trabajo.xlsx"))
    year = int(os.getenv("WORKPLAN_YEAR", "2026"))
    today = date.fromisoformat(os.getenv("WORKPLAN_TODAY", date.today().isoformat()))

    if not token:
        raise RuntimeError("GITHUB_TOKEN is required. Put it in .env or export it before running.")

    gh = GitHub(token)
    items = load_workplan(excel_path)
    repo = get_repo(gh, repository)
    project = ensure_project(gh, owner, repo["node_id"], project_title)

    all_roles = {f"rol: {role}" for item in items for role in split_roles(item.responsable)}
    all_sprints = {f"sprint: {month:02d}-{MONTH_NAMES[month].lower()}" for item in items for month in item.month_numbers}
    ensure_labels(gh, repository, all_roles | all_sprints | {"workplan"})
    milestones = ensure_milestones(gh, repository, items, year)

    fields = get_project_fields(gh, project["id"])
    roles = sorted({role for item in items for role in split_roles(item.responsable)})
    sprints = [f"Sprint {month:02d} - {MONTH_NAMES[month]} {year}" for month in range(1, 13)]
    role_field = ensure_field(gh, project["id"], fields, "Rol", "SINGLE_SELECT", roles)
    sprint_field = ensure_field(gh, project["id"], fields, "Sprint", "SINGLE_SELECT", sprints)
    start_field = ensure_field(gh, project["id"], fields, "Inicio", "DATE")
    end_field = ensure_field(gh, project["id"], fields, "Fin", "DATE")
    status_field = fields.get("Status")
    fields = get_project_fields(gh, project["id"])
    ensure_standard_views(gh, project["id"], fields)

    existing = find_existing_issues(gh, repository)
    created_or_updated = 0
    for item in items:
        issue = create_or_update_issue(gh, repository, item, year, milestones, existing)
        project_item_id = add_issue_to_project(gh, project["id"], issue["node_id"])
        if item.month_numbers:
            sprint_name = f"Sprint {min(item.month_numbers):02d} - {MONTH_NAMES[min(item.month_numbers)]} {year}"
            option_id = single_select_option_id(sprint_field, sprint_name)
            if option_id:
                set_project_field(gh, project["id"], project_item_id, sprint_field["id"], {"singleSelectOptionId": option_id})
        if status_field:
            option_id = single_select_option_id(status_field, status_from_item(item, today, year))
            if option_id:
                set_project_field(gh, project["id"], project_item_id, status_field["id"], {"singleSelectOptionId": option_id})
        first_role = split_roles(item.responsable)[0] if split_roles(item.responsable) else ""
        option_id = single_select_option_id(role_field, first_role)
        if option_id:
            set_project_field(gh, project["id"], project_item_id, role_field["id"], {"singleSelectOptionId": option_id})
        start, end = full_month_range(item.month_numbers, year)
        if start:
            set_project_field(gh, project["id"], project_item_id, start_field["id"], {"date": start})
        if end:
            set_project_field(gh, project["id"], project_item_id, end_field["id"], {"date": end})
        created_or_updated += 1

    print(f"Project: {project['title']} {project['url']}")
    print(f"Repository: {repository}")
    print(f"Work items synced: {created_or_updated}")
    print(f"Milestones/products: {len(milestones)}")
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except Exception as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        raise SystemExit(1)
